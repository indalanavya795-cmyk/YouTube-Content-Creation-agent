import streamlit as st
import json
from datetime import datetime
from io import BytesIO

from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document

from agent import (
    generate_youtube_idea,
    generate_youtube_titles,
    generate_youtube_description,
    generate_youtube_hashtags,
    generate_youtube_keywords,
    generate_youtube_script,
    generate_thumbnail_ideas,
    generate_scene_by_scene_script,
    generate_youtube_shorts,
    generate_instagram_reel,
    generate_repurposed_content,
    generate_seo_analysis,
    save_youtube_content
)


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="YouTube Content Creation Agent",
    page_icon="🎬",
    layout="wide",
    initial_sidebar_state="expanded"
)


# ============================================================
# HISTORY FUNCTIONS
# ============================================================

HISTORY_FILE = "data/content_history.json"


def load_history():
    try:
        with open(HISTORY_FILE, "r", encoding="utf-8") as file:
            return json.load(file)
    except (FileNotFoundError, json.JSONDecodeError):
        return []


def save_history(project):
    history = load_history()
    history.insert(0, project)

    with open(HISTORY_FILE, "w", encoding="utf-8") as file:
        json.dump(
            history,
            file,
            indent=4,
            ensure_ascii=False
        )


# ============================================================
# PDF EXPORT
# ============================================================

def create_pdf(content):
    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()
    story = []

    for section in content.split("\n\n"):
        if section.strip():

            text = section.replace("&", "&amp;")
            text = text.replace("<", "&lt;")
            text = text.replace(">", "&gt;")
            text = text.replace("\n", "<br/>")

            story.append(
                Paragraph(
                    text,
                    styles["BodyText"]
                )
            )

            story.append(
                Spacer(1, 10)
            )

    document.build(story)

    buffer.seek(0)

    return buffer.getvalue()


# ============================================================
# DOCX EXPORT
# ============================================================

def create_docx(content):
    document = Document()

    document.add_heading(
        "YouTube Content Creation Agent",
        level=1
    )

    for section in content.split("\n\n"):
        if section.strip():
            document.add_paragraph(section)

    buffer = BytesIO()

    document.save(buffer)

    buffer.seek(0)

    return buffer.getvalue()


# ============================================================
# SESSION STATE
# ============================================================

default_states = {
    "ideas": "",
    "selected_idea": "",
    "titles": "",
    "description": "",
    "hashtags": "",
    "keywords": "",
    "script": "",
    "thumbnail_ideas": "",
    "scene_plan": "",
    "shorts": "",
    "reel": "",
    "repurposed": "",
    "seo": ""
}


for key, value in default_states.items():
    if key not in st.session_state:
        st.session_state[key] = value


# ============================================================
# HEADER
# ============================================================

st.title("🎬 YouTube Content Creation Agent")

st.markdown(
    """
    ### Create complete YouTube content with AI

    Generate video ideas, titles, descriptions, scripts, SEO content,
    thumbnails ideas and social-media content from one place.
    """
)

st.divider()


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.header("⚙️ Video Settings")

    topic = st.text_input(
        "📌 YouTube Topic",
        placeholder="Example: AI tools for students"
    )

    audience = st.selectbox(
        "🎯 Target Audience",
        [
            "General Audience",
            "College Students",
            "Beginners",
            "Professionals",
            "Content Creators",
            "Entrepreneurs"
        ]
    )

    tone = st.selectbox(
        "🎭 Content Style",
        [
            "Informative",
            "Professional",
            "Friendly",
            "Funny",
            "Storytelling",
            "Energetic"
        ]
    )

    video_length = st.selectbox(
        "⏱️ Video Length",
        [
            "Short (1–3 minutes)",
            "Medium (5–8 minutes)",
            "Long (10–15 minutes)"
        ]
    )

    language = st.selectbox(
        "🌐 Language",
        [
            "English",
            "Telugu",
            "Hindi",
            "Tamil",
            "Kannada"
        ]
    )

    st.divider()

    st.info(
        "💡 Tip: Choose a specific topic for better AI-generated content."
    )


# ============================================================
# VIDEO IDEA GENERATION
# ============================================================

st.header("💡 Generate Video Ideas")

if st.button(
    "✨ Generate Video Ideas",
    use_container_width=True
):

    if not topic.strip():

        st.warning(
            "Please enter a YouTube topic in the sidebar."
        )

    else:

        with st.spinner("Generating creative video ideas..."):

            st.session_state["ideas"] = generate_youtube_idea(
                topic,
                audience,
                tone,
                video_length,
                language
            )

        st.success("Video ideas generated!")


if st.session_state["ideas"]:

    st.subheader("🎯 Generated Ideas")

    st.code(
        st.session_state["ideas"],
        language=None
    )

    selected_idea = st.text_area(
        "✏️ Enter the video idea you want to use",
        value=st.session_state["selected_idea"],
        height=100
    )

    st.session_state["selected_idea"] = selected_idea


# ============================================================
# COMPLETE CONTENT GENERATION
# ============================================================

st.divider()

st.header("🎬 Generate Complete YouTube Content")

if st.button(
    "🚀 Generate Complete Content",
    use_container_width=True
):

    if not st.session_state["selected_idea"].strip():

        st.warning(
            "Please enter or select a video idea first."
        )

    else:

        idea = st.session_state["selected_idea"]

        with st.spinner("Creating your complete YouTube content..."):

            st.session_state["titles"] = generate_youtube_titles(
                idea,
                audience,
                tone,
                video_length,
                language
            )

            st.session_state["description"] = generate_youtube_description(
                idea,
                audience,
                tone,
                video_length,
                language
            )

            st.session_state["hashtags"] = generate_youtube_hashtags(
                idea,
                audience,
                tone,
                video_length,
                language
            )

            st.session_state["keywords"] = generate_youtube_keywords(
                idea,
                audience,
                tone,
                video_length,
                language
            )

            st.session_state["script"] = generate_youtube_script(
                idea,
                audience,
                tone,
                video_length,
                language
            )

            st.session_state["thumbnail_ideas"] = generate_thumbnail_ideas(
                idea,
                audience,
                tone,
                video_length,
                language
            )

        st.success(
            "Complete YouTube content generated successfully!"
        )


# ============================================================
# DISPLAY COMPLETE CONTENT
# ============================================================

if st.session_state["titles"]:

    st.divider()

    st.header("📦 Generated YouTube Content")

    with st.expander("🎯 Titles", expanded=True):

        st.code(
            st.session_state["titles"],
            language=None
        )

    with st.expander("📄 Description"):

        st.code(
            st.session_state["description"],
            language=None
        )

    with st.expander("#️⃣ Hashtags"):

        st.code(
            st.session_state["hashtags"],
            language=None
        )

    with st.expander("🔎 Keywords"):

        st.code(
            st.session_state["keywords"],
            language=None
        )

    with st.expander("🎬 Full Script"):

        st.code(
            st.session_state["script"],
            language=None
        )

    with st.expander("🖼️ Thumbnail Ideas"):

        st.code(
            st.session_state["thumbnail_ideas"],
            language=None
        )


# ============================================================
# ADDITIONAL CONTENT TOOLS
# ============================================================

st.divider()

st.header("🛠️ Additional Content Tools")


# ------------------------------------------------------------
# SCENE BY SCENE
# ------------------------------------------------------------

if st.button(
    "🎞️ Generate Scene-by-Scene Plan",
    use_container_width=True
):

    if not st.session_state["selected_idea"].strip():

        st.warning(
            "Please select a video idea first."
        )

    else:

        with st.spinner("Creating scene-by-scene plan..."):

            st.session_state["scene_plan"] = (
                generate_scene_by_scene_script(
                    st.session_state["selected_idea"],
                    audience,
                    tone,
                    video_length,
                    language
                )
            )

        st.success("Scene-by-scene plan generated!")


if st.session_state["scene_plan"]:

    with st.expander(
        "🎞️ Scene-by-Scene Plan",
        expanded=True
    ):

        st.code(
            st.session_state["scene_plan"],
            language=None
        )


# ------------------------------------------------------------
# YOUTUBE SHORTS
# ------------------------------------------------------------

if st.button(
    "⚡ Generate YouTube Short",
    use_container_width=True
):

    if not st.session_state["selected_idea"].strip():

        st.warning(
            "Please select a video idea first."
        )

    else:

        with st.spinner("Creating YouTube Short..."):

            st.session_state["shorts"] = generate_youtube_shorts(
                st.session_state["selected_idea"],
                audience,
                tone,
                language
            )

        st.success("YouTube Short generated!")


if st.session_state["shorts"]:

    with st.expander(
        "⚡ YouTube Short",
        expanded=True
    ):

        st.code(
            st.session_state["shorts"],
            language=None
        )


# ------------------------------------------------------------
# INSTAGRAM REEL
# ------------------------------------------------------------

if st.button(
    "📱 Generate Instagram Reel",
    use_container_width=True
):

    if not st.session_state["selected_idea"].strip():

        st.warning(
            "Please select a video idea first."
        )

    else:

        with st.spinner("Creating Instagram Reel..."):

            st.session_state["reel"] = generate_instagram_reel(
                st.session_state["selected_idea"],
                audience,
                tone,
                language
            )

        st.success("Instagram Reel generated!")


if st.session_state["reel"]:

    with st.expander(
        "📱 Instagram Reel",
        expanded=True
    ):

        st.code(
            st.session_state["reel"],
            language=None
        )


# ------------------------------------------------------------
# CONTENT REPURPOSING
# ------------------------------------------------------------

if st.button(
    "🔄 Repurpose Content",
    use_container_width=True
):

    if not st.session_state["selected_idea"].strip():

        st.warning(
            "Please select a video idea first."
        )

    elif not st.session_state["script"].strip():

        st.warning(
            "Generate the full script first."
        )

    else:

        with st.spinner("Repurposing your content..."):

            st.session_state["repurposed"] = (
                generate_repurposed_content(
                    st.session_state["selected_idea"],
                    st.session_state["script"],
                    audience,
                    tone,
                    language
                )
            )

        st.success("Content repurposed!")


if st.session_state["repurposed"]:

    with st.expander(
        "🔄 Repurposed Content",
        expanded=True
    ):

        st.code(
            st.session_state["repurposed"],
            language=None
        )


# ------------------------------------------------------------
# SEO ANALYSIS
# ------------------------------------------------------------

if st.button(
    "🔎 Analyze SEO",
    use_container_width=True
):

    if not st.session_state["selected_idea"].strip():

        st.warning(
            "Please select a video idea first."
        )

    else:

        with st.spinner("Analyzing YouTube SEO..."):

            st.session_state["seo"] = generate_seo_analysis(
                st.session_state["selected_idea"],
                st.session_state["titles"],
                st.session_state["keywords"],
                st.session_state["description"]
            )

        st.success("SEO analysis completed!")


if st.session_state["seo"]:

    with st.expander(
        "🔎 SEO Analysis",
        expanded=True
    ):

        st.code(
            st.session_state["seo"],
            language=None
        )


# ============================================================
# SAVE CURRENT CONTENT
# ============================================================

st.divider()

st.header("💾 Save Content")

if st.button(
    "💾 Save Current Project",
    use_container_width=True
):

    if not st.session_state["selected_idea"].strip():

        st.warning(
            "There is no project to save yet."
        )

    else:

        project = {
            "date": datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
            "topic": topic,
            "audience": audience,
            "tone": tone,
            "video_length": video_length,
            "language": language,
            "idea": st.session_state["selected_idea"],
            "titles": st.session_state["titles"],
            "description": st.session_state["description"],
            "hashtags": st.session_state["hashtags"],
            "keywords": st.session_state["keywords"],
            "script": st.session_state["script"],
            "thumbnail_ideas": st.session_state["thumbnail_ideas"],
            "scene_plan": st.session_state["scene_plan"],
            "shorts": st.session_state["shorts"],
            "reel": st.session_state["reel"],
            "repurposed": st.session_state["repurposed"],
            "seo": st.session_state["seo"]
        }

        save_history(project)

        st.success(
            "Project saved to Content History!"
        )


# ============================================================
# CONTENT HISTORY
# ============================================================

st.divider()

st.header("📚 Content History")

history = load_history()

if not history:

    st.info(
        "No saved projects yet."
    )

else:

    project_names = []

    for index, project in enumerate(history):

        project_names.append(
            f"{index + 1}. "
            f"{project.get('idea', 'Untitled')[:70]} "
            f"({project.get('date', '')})"
        )

    selected_project = st.selectbox(
        "Select a previous project",
        project_names
    )

    selected_index = project_names.index(
        selected_project
    )

    previous_project = history[selected_index]

    st.subheader("📖 Previous Project")

    st.write(
        f"**Date:** {previous_project.get('date', '')}"
    )

    st.write(
        f"**Topic:** {previous_project.get('topic', '')}"
    )

    st.write(
        f"**Audience:** {previous_project.get('audience', '')}"
    )

    st.write(
        f"**Style:** {previous_project.get('tone', '')}"
    )

    st.write(
        f"**Language:** {previous_project.get('language', '')}"
    )

    with st.expander("💡 Video Idea"):

        st.code(
            previous_project.get("idea", ""),
            language=None
        )

    with st.expander("🎯 Titles"):

        st.code(
            previous_project.get("titles", ""),
            language=None
        )

    with st.expander("📄 Description"):

        st.code(
            previous_project.get("description", ""),
            language=None
        )

    with st.expander("#️⃣ Hashtags"):

        st.code(
            previous_project.get("hashtags", ""),
            language=None
        )

    with st.expander("🔎 Keywords"):

        st.code(
            previous_project.get("keywords", ""),
            language=None
        )

    with st.expander("🎬 Script"):

        st.code(
            previous_project.get("script", ""),
            language=None
        )

    with st.expander("🖼️ Thumbnail Ideas"):

        st.code(
            previous_project.get(
                "thumbnail_ideas",
                ""
            ),
            language=None
        )

    with st.expander("🎞️ Scene Plan"):

        st.code(
            previous_project.get(
                "scene_plan",
                ""
            ),
            language=None
        )

    with st.expander("⚡ YouTube Short"):

        st.code(
            previous_project.get(
                "shorts",
                ""
            ),
            language=None
        )

    with st.expander("📱 Instagram Reel"):

        st.code(
            previous_project.get(
                "reel",
                ""
            ),
            language=None
        )

    with st.expander("🔄 Repurposed Content"):

        st.code(
            previous_project.get(
                "repurposed",
                ""
            ),
            language=None
        )

    with st.expander("🔎 SEO Analysis"):

        st.code(
            previous_project.get(
                "seo",
                ""
            ),
            language=None
        )


# ============================================================
# EXPORT
# ============================================================

st.divider()

st.header("📤 Export Your Content")

export_content = f"""
VIDEO IDEA

{st.session_state["selected_idea"]}

TITLES

{st.session_state["titles"]}

DESCRIPTION

{st.session_state["description"]}

HASHTAGS

{st.session_state["hashtags"]}

KEYWORDS

{st.session_state["keywords"]}

FULL SCRIPT

{st.session_state["script"]}

THUMBNAIL IDEAS

{st.session_state["thumbnail_ideas"]}

SCENE-BY-SCENE PLAN

{st.session_state["scene_plan"]}

YOUTUBE SHORT

{st.session_state["shorts"]}

INSTAGRAM REEL

{st.session_state["reel"]}

REPURPOSED CONTENT

{st.session_state["repurposed"]}

SEO ANALYSIS

{st.session_state["seo"]}
"""


col1, col2 = st.columns(2)


with col1:

    st.download_button(
        "📄 Download as PDF",
        data=create_pdf(export_content),
        file_name="youtube_content.pdf",
        mime="application/pdf",
        use_container_width=True
    )


with col2:

    st.download_button(
        "📝 Download as DOCX",
        data=create_docx(export_content),
        file_name="youtube_content.docx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        use_container_width=True
    )


# ============================================================
# COPY CONTENT
# ============================================================

st.subheader("📋 Copy Content")

st.write(
    "Use the copy icon in the top-right corner of each box."
)


if st.session_state["titles"]:

    st.write("🎯 Titles")

    st.code(
        st.session_state["titles"],
        language=None
    )


if st.session_state["description"]:

    st.write("📄 Description")

    st.code(
        st.session_state["description"],
        language=None
    )


if st.session_state["hashtags"]:

    st.write("#️⃣ Hashtags")

    st.code(
        st.session_state["hashtags"],
        language=None
    )


if st.session_state["keywords"]:

    st.write("🔎 Keywords")

    st.code(
        st.session_state["keywords"],
        language=None
    )


if st.session_state["script"]:

    st.write("🎬 Script")

    st.code(
        st.session_state["script"],
        language=None
    )


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.caption(
    "🎬 YouTube Content Creation Agent | "
    "Powered by Python, Streamlit and Ollama"
)